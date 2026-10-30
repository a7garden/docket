"""docket 산출물용 엑셀(xlsx) 도구. 템플릿 사본의 칸 값만 바꾸고 나머지는 바이트까지 그대로 둔다.

  inspect <xlsx> [시트!A1:H40]        시트마다 크기 · 병합 · 엑셀 표 · 그림 · 값이 든 칸(수식은 수식 그대로)을 보인다. 범위를 주면 그곳만
  fill <template> <spec.json> <out>   템플릿 사본에 spec 의 작업을 적용해 out 에 쓰고, 수식 계산값을 넣는다. 템플릿은 건드리지 않는다
  new <spec.json> <out>               템플릿이 없을 때 시트 목록으로 새 통합 문서를 만든다
  verify <template> <out>             시트 · 병합 · 열 너비 · 표시 형식 · 그림이 템플릿과 같은지 견주고,
                                      견본 수치가 남은 칸 · 숫자 칸에 들어간 글자를 보인다. 서식이 다르면 종료 코드 1

fill 은 openpyxl 로 다시 저장하지 않는다. 시트 XML 의 칸(<c>)만 고치고 다른 파트는 그대로 옮겨 쓴다 —
openpyxl 로 저장하면 로고 · 그림 · 차트 · 피벗 · 메모 · 인쇄 설정이 사라지는데, 공공 산출물 엑셀은 갑지에 로고가 흔하다.
줄바꿈 칸에 넣은 글이 행 높이를 넘으면 행 높이를 늘린다(엑셀은 연 뒤에 행 높이를 스스로 맞추지 않는다).
수식 계산값은 LibreOffice 로 계산해 넣는다. 비면 docket · 미리 보기 · 지식베이스가 그 칸을 빈칸으로 읽는다.
"""

import copy
import datetime as dt
import json
import math
import os
import re
import shutil
import subprocess
import sys
import tempfile
import unicodedata
import zipfile

from lxml import etree
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils import column_index_from_string, get_column_letter, range_boundaries
from openpyxl.utils.cell import coordinate_from_string

SHOW_MAX = 400
LEFTOVER_MAX = 40
NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PKG_REL = "http://schemas.openxmlformats.org/package/2006/relationships"
OBJECTS = {"xl/drawings/": "그림 · 도형", "xl/charts/": "차트", "xl/pivotTables/": "피벗", "xl/vbaProject.bin": "매크로", "xl/comments": "메모"}
DATE_FMT = "yyyy-mm-dd"
DATETIME_FMT = "yyyy-mm-dd hh:mm"


def q(tag):
	return f"{{{NS}}}{tag}"


def objects_in(path):
	with zipfile.ZipFile(path) as z:
		names = z.namelist()
	return sorted({label for prefix, label in OBJECTS.items() if any(n.startswith(prefix) for n in names)})


def split_ref(ref):
	"""시트!A1 → (시트, A1). 시트 이름에 ! 가 있을 수 있어 마지막 ! 로 가른다."""
	if "!" not in ref:
		raise SystemExit(f"시트 이름이 없다: {ref} (예: 집계!B3)")
	sheet, cell = ref.rsplit("!", 1)
	return sheet.strip("'"), cell.replace("$", "")


def value_of(v):
	"""JSON 값 → 칸 값. {"date": "2026-10-29"} · {"date": "2026-10-29 14:30"} 는 날짜가 된다."""
	if isinstance(v, dict) and "date" in v:
		s = v["date"]
		return dt.datetime.fromisoformat(s) if len(s) > 10 else dt.date.fromisoformat(s)
	return v


def excel_serial(v):
	"""엑셀 날짜 일련번호(1900 체계)."""
	d = v if isinstance(v, dt.datetime) else dt.datetime(v.year, v.month, v.day)
	return (d - dt.datetime(1899, 12, 30)).total_seconds() / 86400


def num_text(v):
	return repr(v) if isinstance(v, float) else str(v)


# ---------- 읽기 ----------


def cmd_inspect(path, rng=None):
	objs = objects_in(path)
	if objs:
		print(f"개체: {' · '.join(objs)} — fill 은 칸 값만 바꾸므로 그대로 남는다")
	wb = load_workbook(path)
	only, box = None, None
	if rng:
		only, area = split_ref(rng)
		box = range_boundaries(area)
	for ws in wb.worksheets:
		if only and ws.title != only:
			continue
		print(f"[{ws.title}] 범위 {ws.dimensions} · 병합 {len(ws.merged_cells.ranges)}개 · 틀 고정 {ws.freeze_panes or '없음'}")
		for r in ws.merged_cells.ranges:
			print(f"  병합 {r.coord}")
		for t in ws.tables.values():
			print(f"  엑셀 표 「{t.displayName}」 {t.ref}")
		n = 0
		rows = ws.iter_rows(min_col=box[0], min_row=box[1], max_col=box[2], max_row=box[3]) if box else ws.iter_rows()
		for row in rows:
			for c in row:
				if c.value is None:
					continue
				n += 1
				if n <= SHOW_MAX:
					fmt = f"\t[{c.number_format}]" if c.number_format != "General" else ""
					wrap = "\t[줄바꿈]" if c.alignment is not None and c.alignment.wrap_text else ""
					print(f"  {ws.title}!{c.coordinate}\t{c.value!r}{fmt}{wrap}")
		if n > SHOW_MAX:
			print(f"  … 값이 든 칸 {n}개 가운데 {SHOW_MAX}개만 보였다. 범위를 주어 나눠 본다")


# ---------- 채우기: 패키지를 그대로 두고 칸만 고친다 ----------


class Package:
	"""xlsx 압축 안 파트. 고친 파트만 다시 쓰고 나머지는 원래 바이트 · 압축 방식 그대로 옮긴다."""

	def __init__(self, path):
		with zipfile.ZipFile(path) as z:
			self.infos = z.infolist()
			self.data = {i.filename: z.read(i.filename) for i in self.infos}
		self.trees = {}
		self.dirty = set()
		self.dropped = set()

	def xml(self, name, write=False):
		"""파트를 읽는다. 고칠 파트는 write=True 로 받는다 — 읽기만 한 파트까지 다시 쓰면 바이트가 바뀌어 원본 보존을 견줄 수 없다."""
		if name not in self.trees:
			self.trees[name] = etree.fromstring(self.data[name])
		if write:
			self.dirty.add(name)
		return self.trees[name]

	def has(self, name):
		return name in self.data and name not in self.dropped

	def rels(self, part):
		folder, base = part.rsplit("/", 1)
		name = f"{folder}/_rels/{base}.rels"
		return name if self.has(name) else None

	def target(self, part, rel_target):
		if rel_target.startswith("/"):
			return rel_target.lstrip("/")
		parts = part.rsplit("/", 1)[0].split("/")
		for seg in rel_target.split("/"):
			if seg == "..":
				parts.pop()
			elif seg and seg != ".":
				parts.append(seg)
		return "/".join(parts)

	def save(self, out):
		with zipfile.ZipFile(out, "w") as z:
			for info in self.infos:
				if info.filename in self.dropped:
					continue
				data = self.data[info.filename]
				if info.filename in self.dirty:
					data = etree.tostring(self.trees[info.filename], xml_declaration=True, encoding="UTF-8", standalone=True)
				z.writestr(info, data, compress_type=info.compress_type)


class Styles:
	"""cellXfs 에서 줄바꿈 · 글자 크기를 읽고, 날짜 칸에 쓸 서식을 하나 더 만든다(원래 서식을 본떠 표시 형식만 바꿈)."""

	def __init__(self, pkg):
		self.pkg = pkg
		self.name = "xl/styles.xml"
		self.root = pkg.xml(self.name) if pkg.has(self.name) else None
		self.cache = {}

	def xfs(self):
		return self.root.find(q("cellXfs")) if self.root is not None else None

	def xf(self, s):
		xfs = self.xfs()
		items = list(xfs) if xfs is not None else []
		return items[s] if 0 <= s < len(items) else None

	def wrap(self, s):
		x = self.xf(s)
		al = x.find(q("alignment")) if x is not None else None
		return al is not None and al.get("wrapText") in ("1", "true")

	def font_size(self, s):
		x = self.xf(s)
		fonts = self.root.find(q("fonts")) if self.root is not None else None
		if x is None or fonts is None:
			return 11.0
		f = list(fonts)[int(x.get("fontId", "0"))]
		sz = f.find(q("sz"))
		return float(sz.get("val")) if sz is not None else 11.0

	def fmt_code(self, s):
		"""서식 s 의 표시 형식 글. 내장 날짜 형식(14~22)은 대표 꼴로 돌려준다."""
		x = self.xf(s)
		if x is None:
			return ""
		fid = int(x.get("numFmtId", "0"))
		if 14 <= fid <= 17:
			return "m/d/yy"
		if 18 <= fid <= 22:
			return "m/d/yy h:mm"
		fmts = self.root.find(q("numFmts"))
		code = next((f.get("formatCode", "") for f in (fmts if fmts is not None else []) if f.get("numFmtId") == str(fid)), "")
		return re.sub(r'"[^"]*"|\[[^\]]*\]', "", code.lower())

	def is_date_fmt(self, s, with_time=False):
		"""날짜 서식인가. with_time 이면 시각(h)까지 보이는 서식이어야 한다 — 날짜만 보이는 칸에 시각을 넣으면 시각이 가려진다."""
		code = self.fmt_code(s)
		return bool(re.search(r"[ymd]", code)) and (not with_time or "h" in code)

	def date_style(self, s, with_time):
		"""날짜 칸 서식. 이미 날짜 서식이면 그대로, 아니면 s 를 본뜬 서식을 새로 붙인다 — 일반 서식 칸에 날짜를 넣으면 46324 로 보인다."""
		if self.root is None or self.is_date_fmt(s, with_time):
			return s
		key = (s, with_time)
		if key in self.cache:
			return self.cache[key]
		self.pkg.xml(self.name, write=True)
		code = DATETIME_FMT if with_time else DATE_FMT
		fmts = self.root.find(q("numFmts"))
		if fmts is None:
			fmts = etree.Element(q("numFmts"))
			self.root.insert(0, fmts)
		fid = next((f.get("numFmtId") for f in fmts if f.get("formatCode") == code), None)
		if fid is None:
			fid = str(max([163] + [int(f.get("numFmtId")) for f in fmts]) + 1)
			etree.SubElement(fmts, q("numFmt"), numFmtId=fid, formatCode=code)
			fmts.set("count", str(len(fmts)))
		xfs = self.xfs()
		base = self.xf(s)
		x = copy.deepcopy(base) if base is not None else etree.Element(q("xf"), numFmtId="0", fontId="0", fillId="0", borderId="0", xfId="0")
		x.set("numFmtId", fid)
		x.set("applyNumberFormat", "1")
		xfs.append(x)
		xfs.set("count", str(len(xfs)))
		self.cache[key] = len(xfs) - 1
		return self.cache[key]


class Sheet:
	def __init__(self, pkg, part, styles):
		self.pkg, self.part, self.styles = pkg, part, styles
		# 시트 객체는 작업이 가리킨 시트에만 만든다. 그래서 만들 때 고칠 파트로 적는다.
		self.root = pkg.xml(part, write=True)
		self.data = self.root.find(q("sheetData"))
		self.rows = {int(r.get("r")): r for r in self.data.iter(q("row"))}
		self.touched = set()
		self.formula_changed = False
		self.merges = [range_boundaries(m.get("ref")) for m in self.root.iter(q("mergeCell"))]

	def row(self, n, like=None):
		r = self.rows.get(n)
		if r is not None:
			return r
		r = etree.Element(q("row"), r=str(n))
		if like is not None and like in self.rows:
			for k in ("ht", "customHeight", "s", "customFormat"):
				if self.rows[like].get(k) is not None:
					r.set(k, self.rows[like].get(k))
		after = [k for k in self.rows if k < n]
		if after:
			self.rows[max(after)].addnext(r)
		else:
			self.data.insert(0, r)
		self.rows[n] = r
		return r

	def cell(self, ref, like_row=None):
		col, n = coordinate_from_string(ref)
		ci = column_index_from_string(col)
		for c0, r0, c1, r1 in self.merges:
			if c0 <= ci <= c1 and r0 <= n <= r1 and (ci, n) != (c0, r0):
				raise SystemExit(f"병합 안쪽 칸이다: {ref} — 병합 첫 칸 {get_column_letter(c0)}{r0} 에 쓴다")
		row = self.row(n, like_row)
		for c in row.iter(q("c")):
			if c.get("r") == ref:
				return c
		c = etree.Element(q("c"), r=ref)
		before = [x for x in row.iter(q("c")) if column_index_from_string(coordinate_from_string(x.get("r"))[0]) < ci]
		if before:
			before[-1].addnext(c)
		else:
			row.insert(0, c)
		if like_row is not None:
			src = self.find(f"{col}{like_row}")
			if src is not None and src.get("s"):
				c.set("s", src.get("s"))
		elif row.get("customFormat") in ("1", "true") and row.get("s"):
			c.set("s", row.get("s"))
		return c

	def find(self, ref):
		col, n = coordinate_from_string(ref)
		row = self.rows.get(n)
		if row is None:
			return None
		return next((c for c in row.iter(q("c")) if c.get("r") == ref), None)

	def put(self, ref, v, like_row=None):
		c = self.cell(ref, like_row)
		f = c.find(q("f"))
		if f is not None and f.get("t") == "shared" and f.get("ref"):
			raise SystemExit(f"공유 수식의 기준 칸이다: {ref} — 이 칸을 바꾸면 같은 수식을 쓰는 다른 칸이 깨진다. 사용자에게 알린다")
		if f is not None or (isinstance(v, str) and v.startswith("=")):
			self.formula_changed = True
		for child in list(c):
			c.remove(child)
		c.attrib.pop("t", None)
		v = value_of(v)
		if v is None:
			pass
		elif isinstance(v, bool):
			c.set("t", "b")
			etree.SubElement(c, q("v")).text = "1" if v else "0"
		elif isinstance(v, (int, float)):
			etree.SubElement(c, q("v")).text = num_text(v)
		elif isinstance(v, (dt.date, dt.datetime)):
			c.set("s", str(self.styles.date_style(int(c.get("s", "0")), isinstance(v, dt.datetime))))
			etree.SubElement(c, q("v")).text = num_text(excel_serial(v))
		elif isinstance(v, str) and v.startswith("="):
			etree.SubElement(c, q("f")).text = v[1:]
		else:
			# 칸 안 글자(inlineStr). 공유 문자열 표를 건드리지 않아도 되고, 엑셀 · LibreOffice · docket 모두 읽는다.
			c.set("t", "inlineStr")
			t = etree.SubElement(etree.SubElement(c, q("is")), q("t"))
			t.text = str(v)
			t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
		self.touched.add(ref)

	def col_width(self, ci):
		cols = self.root.find(q("cols"))
		for col in cols if cols is not None else []:
			if int(col.get("min")) <= ci <= int(col.get("max")) and col.get("width"):
				return float(col.get("width"))
		fmt = self.root.find(q("sheetFormatPr"))
		if fmt is not None and fmt.get("defaultColWidth"):
			return float(fmt.get("defaultColWidth"))
		return 8.43

	def default_height(self):
		fmt = self.root.find(q("sheetFormatPr"))
		return float(fmt.get("defaultRowHeight", "15")) if fmt is not None else 15.0

	def text_of(self, c, shared):
		t = c.get("t")
		if t == "inlineStr":
			return "".join(x.text or "" for x in c.iter(q("t")))
		v = c.find(q("v"))
		if v is None or v.text is None:
			return None
		if t == "s":
			return shared[int(v.text)] if int(v.text) < len(shared) else None
		return v.text if t == "str" else None

	def fit_rows(self, everything=False, shared=()):
		"""줄바꿈 칸의 글이 행 높이를 넘으면 행을 늘린다. 글 폭은 어림한다(한글 · 전각 = 2, 그 밖 = 1, 글자 크기 11pt 기준).
		everything 이면 고친 칸만이 아니라 시트의 모든 글 칸을 본다 — 템플릿에 이미 잘린 행이 흔하다.
		줄이지는 않는다 — 템플릿이 일부러 높게 둔 행일 수 있다."""
		grown = []
		refs = [c.get("r") for c in self.data.iter(q("c"))] if everything else sorted(self.touched)
		for ref in refs:
			c = self.find(ref)
			text = self.text_of(c, shared) if c is not None else None
			if not text:
				continue
			s = int(c.get("s", "0"))
			if not self.styles.wrap(s):
				continue
			col, n = coordinate_from_string(ref)
			ci = column_index_from_string(col)
			c1 = next((m[2] for m in self.merges if m[0] == ci and m[1] == n), ci)
			r1 = next((m[3] for m in self.merges if m[0] == ci and m[1] == n), n)
			if r1 != n:
				continue  # 세로로 병합된 칸은 높이를 여러 행이 나눠 가져 어림하지 않는다
			width = sum(self.col_width(k) for k in range(ci, c1 + 1)) - 1
			size = self.styles.font_size(s)
			lines = 0
			for line in text.split("\n"):
				w = sum(2 if unicodedata.east_asian_width(ch) in "WF" else 1 for ch in line) * size / 11
				lines += max(1, math.ceil(w / max(width, 1)))
			# 한 줄 높이는 글자 크기의 1.35배쯤이다(10pt → 13.5pt · 11pt → 15pt, 엑셀 기본 행 높이와 같다).
			need = round(lines * size * 1.35, 1)
			row = self.rows[n]
			have = float(row.get("ht")) if row.get("ht") else self.default_height()
			if need > have + 0.5:
				row.set("ht", str(need))
				row.set("customHeight", "1")
				grown.append(f"{n}행 {have:g}→{need:g}pt")
		return grown

	def grow_dimension(self):
		dim = self.root.find(q("dimension"))
		if dim is None or not self.touched:
			return
		c0, r0, c1, r1 = range_boundaries(dim.get("ref") if ":" in dim.get("ref", "") else f"{dim.get('ref', 'A1')}:{dim.get('ref', 'A1')}")
		for ref in self.touched:
			col, n = coordinate_from_string(ref)
			ci = column_index_from_string(col)
			c0, r0, c1, r1 = min(c0, ci), min(r0, n), max(c1, ci), max(r1, n)
		dim.set("ref", f"{get_column_letter(c0)}{r0}:{get_column_letter(c1)}{r1}")

	def grow_tables(self, col0, col1, last_row):
		"""엑셀 표 범위가 채운 행을 덮지 못하면 아래로 늘린다. 늘리지 않으면 정렬 · 필터 · 표 서식이 새 행에 안 먹는다."""
		rels = self.pkg.rels(self.part)
		if not rels:
			return
		for rel in self.pkg.xml(rels):
			if not rel.get("Type", "").endswith("/table"):
				continue
			name = self.pkg.target(self.part, rel.get("Target"))
			t = self.pkg.xml(name)
			c0, r0, c1, r1 = range_boundaries(t.get("ref"))
			if c0 <= col0 <= c1 and col1 <= c1 and r0 <= last_row and last_row > r1:
				self.pkg.xml(name, write=True)
				t.set("ref", f"{get_column_letter(c0)}{r0}:{get_column_letter(c1)}{last_row}")
				af = t.find(q("autoFilter"))
				if af is not None:
					af.set("ref", t.get("ref"))
				print(f"엑셀 표 「{t.get('displayName')}」 범위를 {t.get('ref')} 로 늘렸다")


def shared_strings(pkg):
	"""공유 문자열 표. 소리 표기(rPh)는 뺀다."""
	name = "xl/sharedStrings.xml"
	if not pkg.has(name):
		return []
	out = []
	for si in pkg.xml(name).iter(q("si")):
		out.append("".join(t.text or "" for t in si.iter(q("t")) if not any(a.tag == q("rPh") for a in t.iterancestors())))
	return out


def sheet_parts(pkg):
	"""시트 이름 → 압축 안 XML 경로."""
	wb = pkg.xml("xl/workbook.xml")
	rels = pkg.xml("xl/_rels/workbook.xml.rels")
	target = {r.get("Id"): r.get("Target") for r in rels}
	return {s.get("name"): pkg.target("xl/workbook.xml", target[s.get(f"{{{REL_NS}}}id")]) for s in wb.iter(q("sheet"))}


def drop_calc_chain(pkg):
	"""수식을 바꿨으면 계산 순서 기록(calcChain)을 뺀다. 남기면 엑셀이 「복구했다」고 알릴 수 있다 — 엑셀은 없으면 다시 만든다."""
	name = "xl/calcChain.xml"
	if not pkg.has(name):
		return
	pkg.dropped.add(name)
	rels = pkg.xml("xl/_rels/workbook.xml.rels", write=True)
	for r in list(rels):
		if r.get("Target", "").endswith("calcChain.xml"):
			rels.remove(r)
	ct = pkg.xml("[Content_Types].xml", write=True)
	for o in list(ct):
		if o.get("PartName") == "/xl/calcChain.xml":
			ct.remove(o)


def cmd_fill(template, spec_path, out, *flags):
	check_out(template, out)
	pkg = Package(template)
	styles = Styles(pkg)
	parts = sheet_parts(pkg)
	sheets = {}

	def sheet(name):
		if name not in parts:
			raise SystemExit(f"없는 시트: {name} (있는 시트: {', '.join(parts)})")
		if name not in sheets:
			sheets[name] = Sheet(pkg, parts[name], styles)
		return sheets[name]

	spec = json.load(open(spec_path, encoding="utf-8"))
	fit_all = set()
	for op in spec.get("ops", []):
		if "set" in op:
			name, ref = split_ref(op["set"])
			sheet(name).put(ref, op["value"])
		elif "rows" in op:
			# 표 몸통을 채운다. 행을 끼우지 않고 아래로 덮어 쓴다 — 끼우면 병합 · 수식 범위 · 조건부 서식이 따라 밀리지 않아 표가 깨진다.
			name, ref = split_ref(op["rows"])
			sh = sheet(name)
			col, start = coordinate_from_string(ref)
			c0 = column_index_from_string(col)
			like = op.get("style_from")
			width = 0
			for i, values in enumerate(op["values"]):
				width = max(width, len(values))
				for j, v in enumerate(values):
					sh.put(f"{get_column_letter(c0 + j)}{start + i}", v, like)
			sh.grow_tables(c0, c0 + max(width, 1) - 1, start + len(op["values"]) - 1)
		elif "fit_rows" in op:
			fit_all.add(op["fit_rows"])
			sheet(op["fit_rows"])
		elif "clear" in op:
			name, rng = split_ref(op["clear"])
			sh = sheet(name)
			c0, r0, c1, r1 = range_boundaries(rng)
			for n in range(r0, r1 + 1):
				for ci in range(c0, c1 + 1):
					ref = f"{get_column_letter(ci)}{n}"
					if sh.find(ref) is not None:
						sh.put(ref, None)
		else:
			raise SystemExit(f"모르는 작업: {op}")
	changed_formula = False
	shared = shared_strings(pkg) if fit_all else []
	for name, sh in sheets.items():
		for line in sh.fit_rows(name in fit_all, shared):
			print(f"행 높이: [{name}] {line}")
		sh.grow_dimension()
		changed_formula |= sh.formula_changed
	if changed_formula:
		drop_calc_chain(pkg)
	# 엑셀은 열 때 다시 계산하게 한다. 아래에서 넣는 값은 그 전에 읽는 쪽(미리 보기 · docket)을 위한 것이다.
	wbx = pkg.xml("xl/workbook.xml", write=True)
	calc = wbx.find(q("calcPr"))
	if calc is None:
		calc = etree.SubElement(wbx, q("calcPr"))
	calc.set("fullCalcOnLoad", "1")
	pkg.save(out)
	print(f"썼다: {out}")
	bake_values(out)


# ---------- 수식 계산값 ----------


def soffice():
	found = shutil.which("soffice") or shutil.which("soffice.com")
	if found:
		return found
	for base in (os.environ.get("ProgramFiles", ""), os.environ.get("ProgramFiles(x86)", "")):
		p = os.path.join(base, "LibreOffice", "program", "soffice.com")
		if base and os.path.exists(p):
			return p
	return None


def bake_values(out):
	"""LibreOffice 로 계산한 사본에서 수식 칸의 값을 읽어, 결과 파일의 그 칸에 계산값(<v>)만 넣는다.
	사본을 그대로 쓰지 않는 까닭: LibreOffice 로 다시 저장하면 스타일 · 조건부 서식 · 열 너비가 조금씩 바뀐다."""
	probe = Package(out)
	if not any(next(probe.xml(p).iter(q("f")), None) is not None for p in sheet_parts(probe).values()):
		return  # 수식이 없으면 LibreOffice 를 띄우지 않는다(몇 초씩 걸린다)
	exe = soffice()
	if not exe:
		print("계산값: LibreOffice 가 없어 넣지 못했다 — 엑셀은 열 때 계산하지만, 미리 보기 · docket 은 수식 칸을 빈칸으로 읽는다")
		return
	with tempfile.TemporaryDirectory() as tmp:
		src = os.path.join(tmp, "in.xlsx")
		shutil.copy(out, src)
		profile = "file:///" + os.path.join(tmp, "profile").replace("\\", "/")
		r = subprocess.run([exe, f"-env:UserInstallation={profile}", "--headless", "--calc", "--convert-to", "xlsx", "--outdir", os.path.join(tmp, "o"), src], capture_output=True, timeout=180)
		calc = os.path.join(tmp, "o", "in.xlsx")
		if r.returncode != 0 or not os.path.exists(calc):
			print("계산값: LibreOffice 계산이 실패해 넣지 못했다")
			return
		vals = load_workbook(calc, data_only=True)
		pkg = Package(out)
		parts = sheet_parts(pkg)
		n = 0
		for name, part in parts.items():
			if name not in vals.sheetnames:
				continue
			for c in pkg.xml(part).iter(q("c")):
				if c.find(q("f")) is None:
					continue
				v = vals[name][c.get("r")].value
				if v is None:
					continue
				pkg.xml(part, write=True)
				old = c.find(q("v"))
				if old is not None:
					c.remove(old)
				node = etree.SubElement(c, q("v"))
				if isinstance(v, bool):
					c.set("t", "b")
					node.text = "1" if v else "0"
				elif isinstance(v, (int, float)):
					c.attrib.pop("t", None)
					node.text = num_text(v)
				elif isinstance(v, str) and v.startswith("#"):
					c.set("t", "e")
					node.text = v
				elif isinstance(v, (dt.date, dt.datetime)):
					c.attrib.pop("t", None)
					node.text = num_text(excel_serial(v))
				else:
					c.set("t", "str")
					node.text = str(v)
				n += 1
		if n:
			patched = os.path.join(tmp, "patched.xlsx")
			pkg.save(patched)
			shutil.move(patched, out)
		print(f"계산값: 수식 칸 {n}개에 LibreOffice 로 계산한 값을 넣었다")


# ---------- 새로 만들기 · 견주기 ----------


def put_new(cell, v):
	v = value_of(v)
	cell.value = v
	# 날짜를 표시 형식 없는 칸에 넣으면 일련번호(46324)로 보인다.
	if isinstance(v, (dt.date, dt.datetime)) and cell.number_format == "General":
		cell.number_format = DATETIME_FMT if isinstance(v, dt.datetime) else DATE_FMT


def cmd_new(spec_path, out):
	check_out(None, out)
	spec = json.load(open(spec_path, encoding="utf-8"))
	wb = Workbook()
	wb.remove(wb.active)
	for sh in spec.get("sheets", []):
		ws = wb.create_sheet(sh["name"])
		for r in sh.get("rows", []):
			ws.append([value_of(v) for v in r])
		for row in ws.iter_rows():
			for c in row:
				put_new(c, c.value)
		if sh.get("rows") and sh.get("header", True):
			for c in ws[1]:
				c.font = Font(bold=True)
			ws.freeze_panes = "A2"
		for j, w in enumerate(sh.get("widths", []), start=1):
			ws.column_dimensions[get_column_letter(j)].width = w
	wb.save(out)
	print(f"썼다: {out}")
	bake_values(out)


def looks_number(s):
	return bool(re.fullmatch(r"-?[\d,]+(\.\d+)?%?", s.strip()))


def cmd_verify(template, out):
	a, b = load_workbook(template), load_workbook(out)
	bad, same, as_text = [], [], []
	if a.sheetnames != b.sheetnames:
		bad.append(f"시트가 다르다: {a.sheetnames} → {b.sheetnames}")
	for ws in a.worksheets:
		if ws.title not in b.sheetnames:
			continue
		wo = b[ws.title]
		ma, mb = {r.coord for r in ws.merged_cells.ranges}, {r.coord for r in wo.merged_cells.ranges}
		if ma != mb:
			bad.append(f"[{ws.title}] 병합이 다르다: 빠짐 {sorted(ma - mb)} · 생김 {sorted(mb - ma)}")
		wa = {k: v.width for k, v in ws.column_dimensions.items() if v.width}
		wb_ = {k: v.width for k, v in wo.column_dimensions.items() if v.width}
		if wa != wb_:
			bad.append(f"[{ws.title}] 열 너비가 다르다")
		for row in ws.iter_rows():
			for c in row:
				if c.value is None:
					continue
				o = wo[c.coordinate]
				if c.number_format != o.number_format and not isinstance(o.value, (dt.date, dt.datetime)):
					bad.append(f"[{ws.title}] {c.coordinate} 표시 형식 {c.number_format} → {o.number_format}")
				num = isinstance(c.value, (int, float)) and not isinstance(c.value, bool)
				# 글자 이름표는 남는 것이 맞다. 견본 수치가 그대로인 칸이 위험하다.
				if num and o.value == c.value:
					same.append(f"{ws.title}!{c.coordinate}\t{c.value!r}")
				# 숫자 칸에 숫자 모양 글자를 넣으면 합계 · 정렬 · 표시 형식이 안 먹는다. 처음부터 글자인 칸(코드 001 등)은 짚지 않는다.
				if num and isinstance(o.value, str) and looks_number(o.value):
					as_text.append(f"{ws.title}!{c.coordinate}\t{o.value!r}")
	lost = sorted(set(objects_in(template)) - set(objects_in(out)))
	if lost:
		bad.append(f"사라진 개체: {' · '.join(lost)}")
	for line in bad[:50]:
		print("다름:", line)
	print("같다: 시트 · 병합 · 열 너비 · 표시 형식 · 개체" if not bad else f"{len(bad)}건 다르다")
	if as_text:
		print(f"글자가 된 숫자: 템플릿에서 숫자였던 칸에 숫자 모양 글자가 들어갔다 {len(as_text)}곳 — 따옴표 없이 숫자로 넣는다")
		for line in as_text[:LEFTOVER_MAX]:
			print("  " + line)
	if same:
		print(f"견본 수치: 템플릿과 같은 수가 남은 칸 {len(same)}곳 — 이번 자료의 값이 맞는지 본다")
		for line in same[:LEFTOVER_MAX]:
			print("  " + line)
	else:
		print("견본 수치: 템플릿과 같은 수가 남은 칸이 없다")
	sys.exit(1 if bad else 0)


def check_out(template, out):
	if template and os.path.abspath(template) == os.path.abspath(out):
		raise SystemExit("템플릿에 덮어쓰지 않는다. 다른 경로를 준다")
	if os.path.exists(out):
		raise SystemExit(f"이미 있는 파일이다: {out}")


def main(argv):
	if len(argv) < 2 or argv[1] in ("-h", "--help"):
		print(__doc__)
		return
	cmd, args = argv[1], argv[2:]
	{"inspect": cmd_inspect, "fill": cmd_fill, "new": cmd_new, "verify": cmd_verify}[cmd](*args)


if __name__ == "__main__":
	sys.stdout.reconfigure(encoding="utf-8")
	main(sys.argv)
